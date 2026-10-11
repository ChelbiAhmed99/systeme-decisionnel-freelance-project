# generer_soutenance.py — Soutenance PFE TravelBI (Ranim Nabi) — version complète
# Fichiers facultatifs dans le même dossier : logo.png, esprit.png, fond.png, script.md
import math, os, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

# ---------- Charte (codes à ajuster à la pipette sur le logo) ----------
NAVY = RGBColor(0x0B, 0x4F, 0x9E)
ORANGE = RGBColor(0xF3, 0x9A, 0x1E)
SKY = RGBColor(0x2A, 0x9F, 0xD8)
SAND = RGBColor(0xFB, 0xF3, 0xE2)
INK = RGBColor(0x1F, 0x2A, 0x37)
GREY = RGBColor(0x6B, 0x72, 0x80)
GREYT = RGBColor(0x59, 0x59, 0x59)
NOTEG = RGBColor(0xD0, 0xD0, 0xD0)
LIGHT = RGBColor(0xE5, 0xE7, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LOGO, ESPRIT, FOND = "logo.png", "esprit.png", "fond.png"
SHORT = "Système décisionnel et moteur de recommandation de destinations"
RECT, ROUND, OVAL = MSO_SHAPE.RECTANGLE, MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.OVAL
L, C, R = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
FADE = ('<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'spd="med"><p:fade/></p:transition>')

NAV = ["Contexte", "Analyse et conception", "Réalisation",
       "Moteur de recommandation", "Démonstration", "Conclusion"]
TAGSEC = {"Contexte": 0, "Organisme d'accueil": 0, "Étude de l'existant": 0,
          "Critique de l'existant": 0, "Problématique": 0, "Solution proposée": 0,
          "Analyse des besoins": 1, "Conception": 1, "Méthodologie de travail": 1,
          "Réalisation": 2, "Architecture": 2, "Modèles d'IA": 3,
          "Démonstration live": 4, "Diapositive de secours": 4, "Conclusion": 5}
KEY = 0            # numéro de la diapositive dans le script oral
CHROME = set()     # diapositives qui reçoivent la pagination


def load_notes(path="script.md"):
    """Lit le script exporté en Markdown : une section '### N — Titre' par diapositive."""
    notes, cur = {}, None
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            for line in f:
                m = re.match(r"^#{2,4}\s*(\d+)\s*[—–-]\s", line)
                if m:
                    cur = int(m.group(1)); notes[cur] = []
                elif line.startswith("## "):
                    cur = None
                elif cur:
                    notes[cur].append(line.replace("**", "").strip())
    return {k: "\n".join(x for x in v if x) for k, v in notes.items()}


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NOTES = load_notes()


# ---------- Briques de dessin ----------
def fill_text(tf, text, size, color, bold, align, font="Calibri", head=False):
    for i, l in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = l
        r.font.size, r.font.color.rgb, r.font.name = Pt(size), color, font
        r.font.bold = bold or (head and i == 0)


def txt(s, x, y, w, h, text, size=18, color=INK, bold=False, align=L,
        font="Calibri", anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    tb.text_frame.vertical_anchor = anchor
    fill_text(tb.text_frame, text, size, color, bold, align, font)
    return tb


def box(s, x, y, w, h, text="", fill=NAVY, fc=WHITE, size=14, bold=False,
        shape=ROUND, line=None, align=C, head=True, dash=False):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == ROUND:
        sh.adjustments[0] = 0.08
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1.25)
        if dash:
            sh.line.dash_style = MSO_LINE.DASH
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.1)
    fill_text(tf, text, size, fc, bold, align, head=head and "\n" in text)
    return sh


def edge(sh, side):
    x, y, w, h = (v / 914400 for v in (sh.left, sh.top, sh.width, sh.height))
    return {"T": (x + w / 2, y), "B": (x + w / 2, y + h),
            "L": (x, y + h / 2), "R": (x + w, y + h / 2)}[side]


SITE = {"T": 0, "L": 1, "B": 2, "R": 3}


def link(s, a, b, sa="R", sb="L", label="", dash=False, head="arrow",
         color=SKY, width=1.75):
    """Connecteur droit. a et b : formes ou points (x, y) en pouces."""
    p1 = a if isinstance(a, tuple) else edge(a, sa)
    p2 = b if isinstance(b, tuple) else edge(b, sb)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(p1[0]), Inches(p1[1]),
                               Inches(p2[0]), Inches(p2[1]))
    c.line.color.rgb, c.line.width = color, Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    for sh, side, fn in ((a, sa, c.begin_connect), (b, sb, c.end_connect)):
        if not isinstance(sh, tuple) and sh.auto_shape_type != OVAL:
            fn(sh, SITE[side])          # connecteur collé à la forme
    ln = c._element.spPr.get_or_add_ln()
    if head == "arrow":
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    elif head == "diamond":             # composition : losange côté départ
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "diamond", "w": "lg", "len": "lg"}))
    if label:
        txt(s, (p1[0] + p2[0]) / 2 - 0.75, (p1[1] + p2[1]) / 2 - 0.3, 1.5, 0.3,
            label, 11, GREY, align=C)
    return c




def slide(title=None, tag="", extra=False, note=None):
    """extra=True : diapositive ajoutée, garde les notes de la précédente (sauf note=...)."""
    global KEY
    if not extra:
        KEY += 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if title:
        if os.path.exists(FOND):
            s.shapes.add_picture(FOND, 0, 0, width=prs.slide_width, height=prs.slide_height)
        act, w = TAGSEC.get(tag, -1), 12.5 / len(NAV)
        for i, t in enumerate(NAV):                      # chevrons de navigation
            on = i == act
            c = box(s, 0.4 + i * w, 0.15, w + 0.18, 0.62, t, ORANGE if on else WHITE, INK, 11,
                    shape=MSO_SHAPE.CHEVRON, line=None if on else LIGHT)
            c.adjustments[0] = 0.3
            c.shadow.inherit = True
        box(s, 0.5, 0.86, 12.33, 0.02, fill=ORANGE, shape=RECT)
        txt(s, 0.5, 0.88, 11, 0.5, title, 24, GREYT, font="Arial")
        box(s, 0.5, 6.98, 12.33, 0.01, fill=GREY, shape=RECT)
        txt(s, 3.5, 7.03, 6.33, 0.3, SHORT, 10, GREY, align=C)
        if os.path.exists(LOGO):
            s.shapes.add_picture(LOGO, Inches(11.6), Inches(7.02), height=Inches(0.4))
        CHROME.add(s.slide_id)
    s.notes_slide.notes_text_frame.text = note or NOTES.get(
        KEY, f"Coller ici le discours de la diapositive {KEY} (voir le script oral).")
    s._element.append(parse_xml(FADE))
    return s

ANIM = {}          # étapes d'animation par diapositive
P_NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'


def anim(s, steps):
    """steps : liste de groupes de formes ; un groupe apparaît en fondu à chaque clic."""
    ANIM.setdefault(s.slide_id, []).extend([g for g in steps if g])


def timing_xml(steps):
    n = [2]

    def nid():
        n[0] += 1
        return n[0]

    clicks = []
    for group in steps:
        effs = []
        for j, sh in enumerate(group):
            sid, node = sh.shape_id, "clickEffect" if j == 0 else "withEffect"
            effs.append(
                f'<p:par><p:cTn id="{nid()}" presetID="10" presetClass="entr" presetSubtype="0" '
                f'fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                f'<p:childTnLst><p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>'
                f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{nid()}" dur="500"/>'
                f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                f'</p:childTnLst></p:cTn></p:par>')
        clicks.append(
            f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
            f'<p:childTnLst><p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            f'<p:childTnLst>{"".join(effs)}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    return (
        f'<p:timing {P_NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
        f'<p:childTnLst>{"".join(clicks)}</p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')


def cards(s, items, y=1.6, h=1.6, cols=None, fill=SAND, fc=INK, size=16, gap=0.25,
          x0=0.5, width=12.33, animate=True):
    cols = cols or len(items)
    w = (width - gap * (cols - 1)) / cols
    out = []
    for i, t in enumerate(items):
        r, c = divmod(i, cols)
        f = fill[i] if isinstance(fill, list) else fill
        col = WHITE if f in (NAVY, SKY, ORANGE) else fc
        out.append(box(s, x0 + c * (w + gap), y + r * (h + gap), w, h, t, f, col, size))
    if animate:
        anim(s, [[b] for b in out])            # une carte par clic
    return out


def flow(s, items, y, h, fill=SAND, size=14, gap=0.45):
    bs = cards(s, items, y, h, fill=fill, size=size, gap=gap, animate=False)
    steps = [[bs[0]]]
    for a, b in zip(bs, bs[1:]):
        steps.append([link(s, a, b), b])        # flèche + étape suivante
    anim(s, steps)
    return bs


def sticky(s, x, y, w, h, text, fill, rot, mark="pin"):
    """Post-it incliné, avec punaise ou coche ; apparaît au clic."""
    sh = box(s, x, y, w, h, text, fill, INK, 16, shape=RECT)
    sh.shadow.inherit = True
    sh.rotation = rot
    if mark == "pin":
        m = box(s, x + w / 2 - 0.13, y + 0.08, 0.26, 0.26, fill=GREYT, shape=OVAL)
        m.shadow.inherit = True
    else:
        m = txt(s, x + w - 0.75, y + h - 0.85, 0.7, 0.8, "✓", 40, GREYT, True)
    anim(s, [[sh, m]])
    return sh

def actor(s, cx, cy, name):
    """Bonhomme UML centré en (cx, cy)."""
    box(s, cx - 0.16, cy - 0.62, 0.32, 0.32, fill=WHITE, shape=OVAL, line=NAVY)
    for p, q in (((cx, cy - 0.3), (cx, cy + 0.2)), ((cx - 0.28, cy - 0.1), (cx + 0.28, cy - 0.1)),
                 ((cx, cy + 0.2), (cx - 0.22, cy + 0.55)), ((cx, cy + 0.2), (cx + 0.22, cy + 0.55))):
        link(s, p, q, head=None, color=NAVY, width=1.5)
    txt(s, cx - 1.0, cy + 0.58, 2.0, 0.6, name, 13, INK, True, align=C)


def classbox(s, x, y, w, name, attrs, meths, lh=0.18, size=10.5):
    top = box(s, x, y, w, 0.32, name, NAVY, WHITE, 12, True, shape=RECT)
    ha, hm = len(attrs) * lh + 0.1, len(meths) * lh + 0.1
    box(s, x, y + 0.32, w, ha, "\n".join(attrs), WHITE, INK, size, shape=RECT,
        line=NAVY, align=L, head=False)
    bot = box(s, x, y + 0.32 + ha, w, hm, "\n".join(meths), WHITE, INK, size, shape=RECT,
              line=NAVY, align=L, head=False)
    return top, bot


# ================= 1 — Titre =================
s = slide()
if os.path.exists(ESPRIT):
    s.shapes.add_picture(ESPRIT, Inches(5.5), Inches(0.2), height=Inches(0.95))
txt(s, 0.5, 1.2, 12.33, 0.75,
    "Projet de fin d'études\nPour l'obtention du Diplôme National d'Ingénieur · Spécialité : Data Science",
    15, GREYT, True, align=C)
txt(s, 0.8, 2.0, 11.7, 1.4,
    "Conception et réalisation d'un système décisionnel intégrant un moteur de "
    "recommandation de destinations pour une agence de voyages", 28, INK, True, align=C, font="Arial")
fb = s.shapes.build_freeform(Inches(3.2), Inches(5.5))            # faisceau orange
fb.add_line_segments([(Inches(13.333), Inches(3.6)), (Inches(13.333), Inches(4.9))], close=True)
beam = fb.convert_to_shape()
beam.fill.solid(); beam.fill.fore_color.rgb = ORANGE
beam.line.fill.background()
txt(s, 1.5, 3.9, 6, 0.5, "Présenté par : Ranim Nabi", 20, GREYT, True)
if os.path.exists(LOGO):
    s.shapes.add_picture(LOGO, Inches(0.7), Inches(5.9), height=Inches(1.0))
txt(s, 4.6, 5.95, 4.4, 0.8,
    "Encadrant académique : M. Sirajeddine Miladi\nEncadrant professionnel : M. Kamel Kharrat", 13, GREYT)
txt(s, 9.2, 5.95, 3.8, 0.8, "Président : à compléter\nExaminateur : à compléter", 13, GREYT)
txt(s, 9.2, 6.95, 3.8, 0.35, "Soutenu le : à compléter", 13, GREYT, align=R)

# ================= 2 — Plan =================
s = slide("Plan")
for i, t in enumerate(NAV):
    y = 1.6 + i * 0.85
    box(s, 1.1, y, 11.2, 0.62, t, WHITE if i % 2 == 0 else SAND, INK, 20,
        shape=RECT, line=LIGHT).shadow.inherit = True
    box(s, 0.5, y - 0.06, 0.74, 0.74, str(i + 1), ORANGE, WHITE, 20, True, shape=OVAL)

# ================= 3 — Introduction =================
s = slide("Introduction", "Contexte")
cards(s, ["Piloter l'activité\nChiffre d'affaires · canaux de vente\nannulations · saisonnalité",
          "Personnaliser le conseil\nUne proposition adaptée\nà chaque client"],
      y=1.8, h=2.6, fill=[NAVY, ORANGE], size=20)
box(s, 0.5, 4.9, 12.33, 1.2,
    "La donnée, actif stratégique : l'informatique décisionnelle et l'apprentissage automatique "
    "sont désormais accessibles aux petites structures", SAND, INK, 18)

s = slide("Organisme d'accueil : 2 Go Sunny Tunisia Trip", "Organisme d'accueil")
if os.path.exists("agence.png"):          # photo de l'agence ou carte de Monastir
    pic = s.shapes.add_picture("agence.png", Inches(0.6), Inches(1.6), width=Inches(6.6))
else:
    pic = box(s, 0.6, 1.6, 6.6, 3.6, "Insérer ici une photo de l'agence\nou une carte de Monastir (agence.png)",
              WHITE, GREY, 14, shape=RECT, line=GREY, dash=True, head=False)
k = len(s.shapes)
txt(s, 0.6, 5.6, 0.9, 0.9, "4", 48, ORANGE, True, align=C, font="Georgia")
txt(s, 1.5, 5.6, 2.4, 0.9, "canaux de vente", 18, INK, True, anchor=MSO_ANCHOR.MIDDLE)
txt(s, 4.0, 5.6, 0.9, 0.9, "5", 48, ORANGE, True, align=C, font="Georgia")
txt(s, 4.9, 5.6, 2.6, 0.9, "continents couverts", 18, INK, True, anchor=MSO_ANCHOR.MIDDLE)
chiffres = list(s.shapes)[k:]
infos = ["Monastir\navenue Habib Bourguiba", "Directeur\nM. Adel Nabi",
         "Agence de voyages\nsecteur du tourisme", "Vols · hôtels · forfaits\nconseil au client"]
pills = []
for i, t in enumerate(infos):
    y = 1.6 + i * 1.3
    p = box(s, 7.9, y, 4.7, 1.05, t, ORANGE, INK, 16, True, head=False)
    p.adjustments[0] = 0.5
    p.text_frame.margin_left = Inches(1.2)
    ic = box(s, 8.08, y + 0.15, 0.75, 0.75, fill=WHITE, shape=OVAL)     # emplacement de l'icône
    pills.append([p, ic])
anim(s, [[pic]] + pills + [chiffres])

# ================= 5 — Étude de l'existant =================
s = slide("Le système actuel : des tables dispersées", "Étude de l'existant")
tables = [("customers", 0.9, 1.7), ("hotels", 4.0, 1.5), ("flights", 7.3, 1.8), ("flight_price_list", 10.0, 1.5),
          ("packages", 2.3, 3.0), ("bookings", 5.6, 3.1), ("cities", 8.9, 3.2), ("countries", 0.9, 4.2),
          ("continents", 10.6, 4.3)]
for name, x, y in tables:
    box(s, x, y, 2.3, 0.75, name, WHITE, NAVY, 15, True, line=NAVY)
cards(s, ["Suivi de l'activité\nConsultation manuelle des fichiers",
          "Conseil au client\nExpérience du conseiller"], y=5.5, h=1.2, size=17)

# ================= 6 — Critique de l'existant =================
s = slide("Limites du système existant", "Critique de l'existant")
limites = ["Aucun indicateur\nconsolidé", "Saisonnalité difficile\nà repérer", "Pas de segmentation\nformalisée",
           "Conseil sans exploitation\nde l'historique", "Volume réel\nréduit"]
steps = []
for i, t in enumerate(limites):
    x = 0.92 + i * 2.45
    c = box(s, x, 1.6, 1.7, 1.7, str(i + 1), WHITE, LIGHT, 40, True, shape=OVAL, line=ORANGE)
    c.line.width = Pt(7)                                   # anneau ; poser l'icône par-dessus
    lab = txt(s, x - 0.35, 3.4, 2.4, 0.9, t, 15, INK, True, align=C)
    steps.append([c, lab])
k = len(s.shapes)
box(s, 1.5, 4.85, 1.3, 1.15, "!", WHITE, ORANGE, 40, True, shape=MSO_SHAPE.ISOSCELES_TRIANGLE, line=ORANGE)
box(s, 3.3, 4.8, 8.6, 1.2, "", ORANGE, shape=RECT, line=NAVY)
for j, t in enumerate(["Données\ndispersées", "Pilotage sans\nindicateurs fiables", "Conseil dépendant\ndes personnes"]):
    txt(s, 3.4 + j * 2.95, 4.85, 2.5, 1.1, t, 16, INK, True, align=C, anchor=MSO_ANCHOR.MIDDLE)
for j in range(2):
    box(s, 5.95 + j * 2.95, 5.2, 0.45, 0.4, fill=INK, shape=MSO_SHAPE.RIGHT_ARROW)
steps.append(list(s.shapes)[k:])
anim(s, steps)

# ----- ajout : chiffres clés -----
s = slide("L'existant en chiffres", "Critique de l'existant", extra=True,
          note="Diapositive d'appui : laisser les chiffres parler quelques secondes, puis enchaîner sur la problématique.")
kpis = [("11", "clients dans\nl'historique réel"), ("16", "réservations\nréelles"),
        ("9", "tables\nnon reliées"), ("0", "indicateur\nconsolidé"), ("4", "canaux de vente\nnon comparés")]
for i, (n, t) in enumerate(kpis):
    x = 0.5 + i * 2.5
    box(s, x, 2.0, 2.3, 3.4, "", SAND)
    txt(s, x, 2.3, 2.3, 1.4, n, 66, ORANGE, True, align=C, font="Georgia")
    txt(s, x, 4.0, 2.3, 1.2, t, 17, INK, align=C)

# ================= 7 — Problématique =================
s = slide("Problématique", "Problématique")
b = box(s, 0.9, 1.9, 8.4, 4.4,
        "Comment piloter l'activité par des indicateurs fiables et conseiller chaque client "
        "à partir des données de l'agence ?", RGBColor(0xA6, 0xA6, 0xA6), WHITE, 26,
        shape=MSO_SHAPE.OVAL_CALLOUT, line=GREYT)
b.adjustments[0], b.adjustments[1] = 0.65, -0.3          # pointe vers le personnage
box(s, 10.3, 1.9, 2.4, 4.3, "Insérer ici\nune illustration\n(personnage)", WHITE, GREY, 13,
    shape=RECT, line=GREY, dash=True, head=False)

# ================= 8 — Solution proposée =================
s = slide("TravelBI : trois volets, une application", "Solution proposée")
vs = cards(s, ["Volet données\nPipeline Python · modèle en étoile",
               "Volet décisionnel\nRapport Power BI de cinq pages",
               "Volet intelligent\nMoteur de recommandation hybride"],
           y=1.5, h=1.3, fill=[NAVY, SKY, ORANGE], size=16)
app = box(s, 3.2, 3.4, 6.9, 0.8, "Application de bureau TravelBI", WHITE, NAVY, 20, True, line=NAVY)
anim(s, [[app] + [link(s, v, app, "B", "T") for v in vs]])
rows = [("Tables dispersées", "Modèle en étoile"),
        ("Aucun indicateur consolidé", "KPI et tableaux de bord filtrables"),
        ("Expérience du conseiller", "Top-N avec score et explication")]
for i, (a, b) in enumerate(rows):
    y = 4.6 + i * 0.75
    ba = box(s, 1.5, y, 4.3, 0.6, a, LIGHT, INK, 15)
    bb = box(s, 7.5, y, 4.3, 0.6, b, SAND, INK, 15, True)
    anim(s, [[ba, link(s, ba, bb, color=ORANGE), bb]])
# ================= 9 — Besoins fonctionnels =================
s = slide("Besoins fonctionnels", "Analyse des besoins")
cards(s, ["Données · BF1 à BF4\n\nCharger et qualifier\nNettoyer\nIndicateurs client et destination\nSegmenter la clientèle",
          "Décisionnel · BF5, BF6\n\nModèle en étoile et KPI\nTableaux de bord filtrables\n(période, statut familial, sexe)",
          "Recommandation · BF7 à BF9\n\nTop-N expliqué, client existant\nNouveau client par profil\nComparaison des modèles",
          "Application · BF10\n\nIndicateurs, rapports Power BI,\nrecommandations et performances\ndans un seul outil"],
      y=1.6, h=4.6, size=15)
for i, t in enumerate(["Sprint 1", "Sprint 2", "Sprint 3", "Sprint 4"]):
    txt(s, 0.5 + i * 3.145, 6.35, 2.9, 0.4, t, 14, ORANGE, True, align=C)

# ================= 10 — Besoins non fonctionnels =================
s = slide("Besoins non fonctionnels", "Analyse des besoins")
bnf = ["Fiabilité\nDéfinition unique des indicateurs", "Reproductibilité\nGraine aléatoire fixée",
       "Cohérence\nApplication = notebook", "Explicabilité\nJustification compréhensible",
       "Performance\nMoins d'une seconde", "Autonomie\nHors ligne, sans Python",
       "Ergonomie\nInterface en français", "Confidentialité\nDonnées sur le poste de l'agence"]
cards(s, bnf, y=1.7, h=2.3, cols=4, size=15,
      fill=[NAVY, NAVY, ORANGE, ORANGE, NAVY, ORANGE, NAVY, NAVY])

# ================= 11 — Cas d'utilisation =================
s = slide("Diagramme global de cas d'utilisation", "Conception")
box(s, 3.2, 1.45, 6.95, 5.45, "", WHITE, line=NAVY, shape=RECT)
txt(s, 3.2, 1.47, 6.95, 0.35, "Système TravelBI", 14, NAVY, True, align=C)


def uc(x, y, t):
    return box(s, x, y, 3.05, 0.68, t, SAND, INK, 12, shape=OVAL, line=NAVY, head=False)


left = [uc(3.45, y, t) for y, t in ((1.9, "Consulter le tableau de bord"), (2.72, "Consulter les rapports Power BI"),
                                    (3.54, "Consulter la performance\ndes modèles"), (4.5, "Recommander pour\nun client existant"),
                                    (5.32, "Recommander pour\nun nouveau client"), (6.14, "Exporter les\nrecommandations"))]
right = [uc(6.85, y, t) for y, t in ((1.9, "Préparer les données et\nentraîner les modèles"),
                                     (3.54, "Publier le rapport Power BI"), (5.32, "Vérifier la cohérence"))]
def groupe(dessin):
    k = len(s.shapes)
    dessin()
    return list(s.shapes)[k:]


g = groupe(lambda: (actor(s, 1.5, 2.9, "Directeur"),
                    [link(s, (1.85, 2.8), u, sb="L", head=None, color=NAVY, width=1.25) for u in left[:3]]))
anim(s, [left[:3] + g])
g = groupe(lambda: (actor(s, 1.5, 5.5, "Conseiller\ncommercial"),
                    [link(s, (1.85, 5.4), u, sb="L", head=None, color=NAVY, width=1.25) for u in left[3:]]))
anim(s, [left[3:] + g])
g = groupe(lambda: (actor(s, 11.8, 5.2, "Analyste\nde données"),
                    [link(s, (11.45, 5.1), u, sb="R", head=None, color=NAVY, width=1.25) for u in right]))
anim(s, [right + g])
g = groupe(lambda: (actor(s, 11.8, 2.6, "«système»\nPower BI Service"),
                    [link(s, (11.45, 2.5), u, sb="R", head=None, color=NAVY, width=1.25) for u in (left[1], right[1])]))
anim(s, [g])
# ================= 12 — Diagramme de classes =================
s = slide("Diagramme de classes de TravelBI", "Conception")
ft, fb_ = classbox(s, 5.2, 1.42, 2.95, "FenetrePrincipale",
                   ["– moteur : RecommenderEngine", "– config : dict"], ["+ naviguer(page)", "+ recharger()"])
pages = [("PageTableauBord", ["– annee : int"], ["+ actualiser()"]),
         ("PageRapports", ["– liens : list"], ["+ afficher(lien)", "+ modePresentation()"]),
         ("PageRecommandations", ["– mode : str", "– n : int"], ["+ recommander()", "+ exporterCSV()"]),
         ("PageModeles", ["– comparaison : DataFrame"], ["+ afficher()"]),
         ("PageAPropos", ["– etatFichiers : dict"], ["+ verifierCoherence()", "+ recharger()"])]
pb = {}
for i, (n, a, m) in enumerate(pages):
    t, b = classbox(s, 0.5 + i * 2.5, 2.95, 2.33, n, a, m)
    link(s, fb_, t, "B", "T", head="diamond", color=NAVY, width=1.25)
    pb[n] = b
it, _ = classbox(s, 0.5, 4.8, 3.0, "«module» indicateurs", ["– faits : DataFrame"],
                 ["+ calculerKPI(annee)", "+ caMensuel(annee)", "+ topDestinations(annee)"])
et, _ = classbox(s, 6.2, 4.7, 3.6, "RecommenderEngine",
                 ["– meta : dict", "– booster : Booster", "– vecteurs, P, Q, biais", "– poids : dict"],
                 ["+ charger(dossier)", "+ recommanderClient(client_id, n)", "+ recommanderNouveau(profil, n)",
                  "+ expliquer(client, destination)", "+ selfcheck()"])
link(s, pb["PageTableauBord"], it, "B", "T", "«use»", dash=True, color=GREY, width=1.25)
link(s, pb["PageRecommandations"], et, "B", "T", "«use»", dash=True, color=GREY, width=1.25)
link(s, pb["PageAPropos"], et, "B", "T", "«use»", dash=True, color=GREY, width=1.25)

# ================= 13 — Méthodologie =================
s = slide("Une démarche hybride : Scrum et CRISP-DM", "Méthodologie de travail")
seg = [("Cadrage\nS1–S4", 4, GREY, "Spécification\net conception"),
       ("Sprint 1 · Données\nS5–S9", 5, NAVY, "CRISP-DM 1 à 3\nmétier, données,\npréparation"),
       ("Sprint 2 · Décisionnel\nS10–S14", 5, SKY, "CRISP-DM 6\npremier\ndéploiement"),
       ("Sprint 3 · Recommandation\nS15–S20", 6, ORANGE, "CRISP-DM 4 et 5\nmodélisation,\névaluation"),
       ("Sprint 4 · Application\nS21–S24", 4, NAVY, "CRISP-DM 6\ndéploiement"),
       ("Finalisation\nS25–S26", 2, GREY, "")]
x, u = 0.5, 12.33 / 26
for t, wk, col, sub in seg:
    box(s, x, 1.7, wk * u - 0.06, 1.0, t, col, WHITE, 12, shape=RECT)
    if sub:
        txt(s, x, 2.8, wk * u - 0.06, 0.9, sub, 12, GREY, align=C)
    x += wk * u
txt(s, 0.5, 3.75, 12.33, 0.4, "Rédaction du rapport et tests en parallèle, de S3 à S26", 13, GREY, align=C)
cards(s, ["Product Owner\nDirection de l'agence\net encadrant professionnel",
          "Scrum Master\nEncadrant académique", "Équipe de développement\nL'étudiante",
          "Critère du sprint 3\nHitRate@7 ≥ 0,70"],
      y=4.5, h=1.8, fill=[SAND, SAND, SAND, ORANGE], size=15)

# ----- ajout : cycle CRISP-DM -----
s = slide("CRISP-DM appliqué au projet", "Méthodologie de travail", extra=True)
box(s, 5.37, 2.75, 2.6, 2.6, "CRISP-DM\nsix phases itératives", ORANGE, INK, 17, shape=OVAL)
phases = ["1 · Compréhension\ndu métier", "2 · Compréhension\ndes données", "3 · Préparation\ndes données",
          "4 · Modélisation", "5 · Évaluation", "6 · Déploiement"]
for i, t in enumerate(phases):
    a = math.radians(-90 + i * 60)
    cx, cy = 6.67 + 4.0 * math.cos(a), 4.05 + 2.25 * math.sin(a)
    txt(s, cx - 1.4, cy - 0.4, 2.8, 0.8, t, 15, GREYT, True, align=C, anchor=MSO_ANCHOR.MIDDLE)
    m = math.radians(-60 + i * 60)
    ar = box(s, 6.67 + 2.6 * math.cos(m) - 0.3, 4.05 + 1.75 * math.sin(m) - 0.2, 0.6, 0.4,
             fill=GREY, shape=MSO_SHAPE.RIGHT_ARROW)
    ar.rotation = math.degrees(m) + 90

# ================= 14 — Technologies =================
s = slide("Technologies utilisées", "Réalisation")
tech = [("Données et modèles", "Python · Jupyter sur Google Colab · pandas · NumPy · scikit-learn\n"
                               "LightGBM · XGBoost · CatBoost · Matplotlib · Seaborn", NAVY),
        ("Stockage", "Modèle en étoile exporté en fichiers CSV", SKY),
        ("Restitution", "Power BI Desktop · Power BI Service · mesures DAX", SKY),
        ("Application", "PyQt6 · Qt WebEngine · PyInstaller (exécutable Windows)", ORANGE)]
for i, (k, v, col) in enumerate(tech):
    y = 1.6 + i * 1.3
    box(s, 0.5, y, 3.0, 1.1, k, col, WHITE, 17, True)
    box(s, 3.7, y, 9.13, 1.1, v, SAND, INK, 16, align=L, head=False)

# ================= 15 — Architecture logique (3 étapes) =================
layers = [("Sources", NAVY, "Tables de l'agence :\nclients, destinations,\nhôtels, vols,\npréférences,\nréservations"),
          ("Préparation", NAVY, "Notebook Python :\nqualité, nettoyage,\nRFM, vecteurs de profil,\nscore d'interaction,\nsegmentation KMeans"),
          ("Stockage", SKY, "powerbi_exports :\ntable de faits,\n4 dimensions, 7 KPI\n\nTravelBI_App :\nmodèles et données"),
          ("Modèles", ORANGE, "Filtrage par contenu\nFactorisation matricielle\n(poids nul)\nLightGBM\nPopularité\nCombinaison pondérée"),
          ("Restitution", SKY, "Power BI Desktop\npuis Service\n\nTravelBI :\nindicateurs.py\nmoteur_reco.py\napp.py")]
labels = ["charge", "exporte", "alimente", "lit"]


def archi_logique(nmax, banner, extra):
    s = slide("Architecture logique en couches", "Architecture", extra=extra)
    prev = None
    for i, (h, col, body) in enumerate(layers[:nmax]):
        x = 0.5 + i * 2.53
        box(s, x, 1.55, 2.2, 0.55, h, col, WHITE, 16, True, shape=RECT)
        b = box(s, x, 2.1, 2.2, 3.4, body, SAND, INK, 13, shape=RECT, head=False)
        if prev is not None:
            link(s, prev, b, label=labels[i - 1])
        prev = b
    if banner:
        box(s, 0.5, 5.9, 12.33, 0.85,
            "Le notebook, source unique : Power BI et l'application consomment les fichiers de la même exécution",
            NAVY, WHITE, 17, True)


for k, (nmax, banner) in enumerate([(2, False), (4, False), (5, True)]):
    archi_logique(nmax, banner, extra=k > 0)

# ================= 16 — Architecture physique =================
s = slide("Architecture physique (déploiement)", "Architecture")


def node(x, y, w, h, t):
    box(s, x, y, w, h, "", WHITE, line=NAVY, shape=RECT)
    txt(s, x, y + 0.03, w, 0.4, t, 13, NAVY, True, align=C)


def art(x, y, t, w=2.2, fill=SAND):
    return box(s, x, y, w, 0.9, t, fill, INK, 12, shape=RECT, line=NAVY, head=False)


node(0.5, 1.5, 5.3, 5.3, "«device» Poste d'analyse")
nb = art(0.8, 2.3, "«artifact»\nNotebook (Jupyter / Colab)")
ex = art(3.3, 2.3, "«artifact»\npowerbi_exports (CSV)")
pbd = art(3.3, 3.9, "«component»\nPower BI Desktop")
tapp = art(0.8, 5.5, "«artifact»\nTravelBI_App (models, data)")
node(7.4, 1.5, 5.4, 1.9, "«execution environment» Power BI Service")
rp = art(8.6, 2.2, "«artifact»\nRapport publié", 3.0)
node(7.4, 3.9, 5.4, 2.9, "«device» Poste de l'agence")
exe = art(8.6, 4.5, "«artifact»\nTravelBI.exe", 3.0, WHITE)
loc = art(8.6, 5.7, "«artifact»\nmodels, data, config.json", 3.0)
link(s, nb, ex, "R", "L", "exporte", dash=True, color=GREY)
link(s, nb, tapp, "B", "T", "exporte", dash=True, color=GREY)
link(s, ex, pbd, "B", "T", "importe", dash=True, color=GREY)
link(s, pbd, rp, "R", "L", "publication", color=NAVY)
link(s, tapp, loc, "R", "L", "copie", color=NAVY)
link(s, exe, loc, "B", "T", "lit", dash=True, color=GREY)
link(s, exe, rp, "T", "B", "HTTPS", color=ORANGE, head=None, width=2.25)

# ================= 17 — Des données aux variables =================
s = slide("Des données réelles aux interactions", "Modèles d'IA")
flow(s, ["Historique réel\n11 clients\n16 réservations",
         "Augmentation\nbootstrap, perturbation,\ngénération par segment",
         "Jeu de données\n500 clients · 45 destinations\n9 944 réservations",
         "Nettoyage\n28 doublons retirés\n444 montants écrêtés (IQR)\n9 916 réservations",
         "Score d'interaction\nimplicite de 0 à 5\n4 802 interactions"],
     1.8, 2.4, fill=[ORANGE, SAND, SAND, SAND, NAVY], size=13)
cards(s, ["Score implicite\nsatisfaction 45 % · récence 25 % · montant 20 % · constante 10 %",
          "Segmentation KMeans, k = 5\nâge, revenu, fréquence, montant, récence, taux d'annulation"],
      y=4.8, h=1.5, size=15)

# ================= 18 — Moteur hybride (5 étapes) =================
sig = [("Filtrage par contenu\nsimilarité cosinus · poids 0,2", NAVY),
       ("Factorisation matricielle\nNumPy, SGD · poids 0", LIGHT),
       ("LightGBM\n15 variables · poids 0,5", ORANGE),
       ("Popularité\na priori · poids 0,3", SKY)]


def moteur(nsig, fin, extra):
    s = slide("Le moteur de recommandation hybride", "Modèles d'IA", extra=extra)
    src = box(s, 0.5, 3.0, 2.1, 1.6, "Profil client\net destination\n(20 dimensions)", SAND, INK, 14, head=False)
    bs = []
    for i, (t, col) in enumerate(sig[:nsig]):
        b = box(s, 3.3, 1.6 + i * 1.3, 3.4, 1.05, t, col, INK if col == LIGHT else WHITE, 13)
        link(s, src, b)
        bs.append(b)
    if nsig == 4:
        comb = box(s, 7.4, 3.0, 1.9, 1.6, "Combinaison\npondérée\nbornée à [0, 5]", WHITE, NAVY, 14, True, line=NAVY)
        for b in bs:
            link(s, b, comb)
    if fin:
        exc = box(s, 9.9, 1.9, 2.9, 1.2, "Exclusion des destinations\ndéjà réservées", SAND, INK, 13, head=False)
        top = box(s, 9.9, 4.4, 2.9, 1.4, "Top-N\navec score et explication\nen langage naturel", NAVY, WHITE, 14)
        link(s, comb, exc, "R", "L")
        link(s, exc, top, "B", "T")
        txt(s, 0.5, 6.5, 12.33, 0.4,
            "Nouveau client : le profil déclaré remplace l'historique (démarrage à froid)", 14, GREY, align=C)


for k, (nsig, fin) in enumerate([(1, False), (2, False), (3, False), (4, False), (4, True)]):
    moteur(nsig, fin, extra=k > 0)

# ================= 19 — Résultats =================
s = slide("Sept approches comparées (jeu de test, K = 7)", "Modèles d'IA")
cd = CategoryChartData()
cd.categories = ["Contenu", "Factorisation", "LightGBM", "XGBoost", "CatBoost", "Ensemble", "Hybride"]
cd.add_series("HitRate@7", (0.5852, 0.2244, 0.8196, 0.8116, 0.8257, 0.8096, 0.8497))
cd.add_series("NDCG@7", (0.2343, 0.0770, 0.4427, 0.4100, 0.4330, 0.4253, 0.4607))
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.45),
                        Inches(8.4), Inches(5.4), cd).chart
ch.has_legend = True
ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
pl = ch.plots[0]
pl.has_data_labels = True
pl.data_labels.font.size = Pt(11)
pl.data_labels.number_format, pl.data_labels.number_format_is_linked = "0.00", False
for ser, col in zip(pl.series, (NAVY, SKY)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
hp = pl.series[0].points[6]
hp.format.fill.solid(); hp.format.fill.fore_color.rgb = ORANGE
ch.category_axis.tick_labels.font.size = Pt(12)
ch.value_axis.tick_labels.font.size = Pt(11)
box(s, 9.2, 1.6, 3.63, 2.3, "85 %\nHitRate@7 du modèle hybride\nseuil visé : 70 %", ORANGE, WHITE, 18)
box(s, 9.2, 4.2, 3.63, 2.3,
    "RMSE le plus bas (0,354) :\nla factorisation matricielle,\npourtant dernière au classement",
    SAND, INK, 15, head=False)

# ================= 20 — Lecture critique =================
s = slide("Lecture critique des résultats", "Modèles d'IA")
crit = ["Biais de sélection\nPoids hybrides et booster\nchoisis sur le jeu de test",
        "Fuite d'information\nVariables calculées avant\nle découpage",
        "Pertinence peu sélective\n97,6 % des interactions\nau-dessus du seuil"]
for i, t in enumerate(crit):
    sticky(s, 0.8 + i * 4.1, 1.7 + (i % 2) * 0.5, 3.6, 2.3, t, NOTEG if i % 2 == 0 else ORANGE, (-2, 2, -3)[i])
box(s, 0.5, 5.2, 12.33, 1.3,
    "85 % est une borne haute · données majoritairement générées\n"
    "Ce qui est validé : la démarche et la chaîne technique, rejouable sur des données réelles",
    NAVY, WHITE, 17, head=False)

# ================= 21 — Démonstration =================
s = slide("Démonstration : TravelBI en direct", "Démonstration live")
flow(s, ["1\nTableau de bord\nannée 2026", "2\nRapport Power BI\nintégré et filtré",
         "3\nRecommandation\nclient CLI00001", "4\nNouveau client\npar profil déclaré",
         "5\nContrôle de cohérence\navec le notebook"],
     2.3, 2.2, fill=[NAVY, SKY, ORANGE, ORANGE, NAVY], size=15)
txt(s, 0.5, 5.0, 12.33, 0.5, "Deux parcours : le directeur, puis le conseiller face à un client",
    18, GREY, align=C)

# ================= 22 — Secours (masquée) =================
s = slide("Démonstration : captures de l'application", "Diapositive de secours")
s._element.set("show", "0")
caps = ["Insérer la figure 4.8\nTableau de bord", "Insérer la figure 4.9\nRecommandations CLI00001",
        "Insérer la figure 4.11\nRapport Power BI intégré", "Insérer la figure 4.12\nContrôle de cohérence"]
for i, t in enumerate(caps):
    r, c = divmod(i, 2)
    box(s, 0.5 + c * 6.3, 1.45 + r * 2.75, 6.05, 2.6, t, WHITE, GREY, 16, shape=RECT,
        line=GREY, dash=True, head=False)

# ================= 23 — Conclusion =================
s = slide("Conclusion", "Conclusion")
bilan = ["Sprint 1\nDonnées nettoyées, RFM,\ncinq segments de clients",
         "Sprint 2\nModèle en étoile et rapport\nPower BI de cinq pages",
         "Sprint 3\nModèle hybride retenu\nHitRate@7 85 %",
         "Sprint 4\nApplication TravelBI\ncohérence 40 / 40"]
for i, t in enumerate(bilan):
    sticky(s, 0.7 + i * 3.1, 1.7 + (i % 2) * 1.5, 2.9, 2.3, t,
           NOTEG if i % 2 == 0 else ORANGE, (-3, 2, -2, 3)[i], mark="check")
txt(s, 0.5, 6.1, 12.33, 0.6, "Des données dispersées à un outil d'aide à la décision utilisable par l'agence",
    18, GREYT, True, align=C)

# ================= 24 — Perspectives =================
s = slide("Perspectives", "Conclusion")
court = ["Court terme\nCorriger le protocole\nd'évaluation",
         "Court terme\nHarmoniser les indicateurs\nPower BI et application",
         "Court terme\nMesurer la diversité\ndes recommandations"]
long_ = ["Long terme\nCollecter l'historique réel\net réentraîner",
         "Long terme\nMise en production\nsur un serveur multi-postes",
         "Long terme\nExpérimentation contrôlée\nauprès des clients"]
for i in range(3):
    sticky(s, 0.7 + i * 4.15, 1.7, 3.6, 2.0, court[i], NOTEG, (-3, 2, -2)[i])
    sticky(s, 1.1 + i * 4.15, 4.3, 3.6, 2.0, long_[i], ORANGE, (2, -2, 3)[i])

# ================= 25 — Merci =================
s = slide()
if os.path.exists(LOGO):
    s.shapes.add_picture(LOGO, Inches(4.4), Inches(1.0), height=Inches(1.5))
txt(s, 0.5, 3.0, 12.33, 1.0, "Merci pour votre attention", 44, GREYT, True, align=C, font="Arial")
txt(s, 0.5, 5.6, 12.33, 0.9,
    "Ranim Nabi · Encadrants : M. Sirajeddine Miladi et M. Kamel Kharrat\n2 Go Sunny Tunisia Trip · ESPRIT 2025-2026",
    15, INK, align=C)

# ----- ajout : Questions -----
s = slide(extra=True, note="Répondre posément ; revenir aux diapositives d'architecture ou de résultats si besoin.")
b = box(s, 3.2, 1.8, 6.9, 3.6, "Questions", ORANGE, INK, 54, True, shape=MSO_SHAPE.OVAL_CALLOUT)
b.adjustments[0], b.adjustments[1] = -0.3, 0.7

# ---------- Pagination ----------
for i, s in enumerate(prs.slides, 1):
    if s.slide_id in CHROME:
        txt(s, 0.5, 7.03, 1.5, 0.3, f"Page {i}", 10, GREY)
for s in prs.slides:
    if s.slide_id in ANIM:
        s._element.append(parse_xml(timing_xml(ANIM[s.slide_id])))
prs.save("Soutenance_TravelBI.pptx")
print("Créé : Soutenance_TravelBI.pptx —", len(prs.slides), "diapositives")